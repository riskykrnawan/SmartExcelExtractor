"""Background threads so the interface never freezes while reading files."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from PySide6.QtCore import QThread, Signal

from core.audit import RunResult
from core.extractor import run_extraction
from core.json_writer import write_json
from core.models import Profile

logger = logging.getLogger(__name__)

#: Cells shown around the result in the mini-grid (5 on each side).
GRID_RADIUS = 5


class RunWorker(QThread):
    """Run the full extraction and write the JSON result file."""

    succeeded = Signal(object, str)  # RunResult, written path
    failed = Signal(str)  # friendly message (details go to the log)

    def __init__(self, profile: Profile, output_dir: Path) -> None:
        super().__init__()
        # Snapshot so edits in other wizard steps cannot change a running job.
        self._profile = profile.model_copy(deep=True)
        self._output_dir = output_dir

    def run(self) -> None:
        try:
            result = run_extraction(self._profile)
            path = write_json(result, self._output_dir)
            self.succeeded.emit(result, str(path))
        except Exception:
            logger.exception("Extraction failed in the UI worker")
            self.failed.emit("Ekstraksi gagal total. Detail lengkap ada di file log.")


class PreviewWorker(QThread):
    """Run one field against a sample file and build the mini-grid snippet."""

    succeeded = Signal(int, dict)  # request id, payload
    failed = Signal(int, str)  # request id, friendly message

    def __init__(self, request_id: int, profile: Profile) -> None:
        super().__init__()
        self._request_id = request_id
        self._profile = profile.model_copy(deep=True)

    def run(self) -> None:
        try:
            result = run_extraction(self._profile)
            payload: dict[str, Any] = {"result": result.to_dict()}
            payload["grid"] = self._build_grid(result)
            self.succeeded.emit(self._request_id, payload)
        except Exception:
            logger.exception("Preview failed")
            self.failed.emit(
                self._request_id,
                "Pratinjau gagal dijalankan. Detail lengkap ada di file log.",
            )

    def _build_grid(self, result: RunResult) -> dict[str, Any] | None:
        """Read a small window of cells around the result for visual checking."""
        entry = next(iter(result.data.values()), None)
        if entry is None:
            return None
        source = entry["source"]
        location = entry["location"]
        focus = location.get("value_address") or location.get("row_keyword_address")
        if not focus or not source["file"] or not source["sheet"]:
            return None
        input_root = result.meta.get("input_root")
        if not input_root:
            return None
        file_path = Path(input_root) / source["folder"] / source["file"]
        try:
            wb = load_workbook(file_path, data_only=True)
        except Exception:  # noqa: BLE001 - preview simply shows no grid
            logger.warning("Cannot reopen %s for the mini grid", file_path)
            return None
        try:
            ws = wb[source["sheet"]]
            center_row, center_col = coordinate_to_tuple(focus)
            row_start = max(1, center_row - GRID_RADIUS)
            col_start = max(1, center_col - GRID_RADIUS)
            row_end = center_row + GRID_RADIUS
            col_end = center_col + GRID_RADIUS
            rows: list[list[str]] = []
            for row in range(row_start, row_end + 1):
                values: list[str] = []
                for col in range(col_start, col_end + 1):
                    value = ws.cell(row=row, column=col).value
                    values.append("" if value is None else str(value))
                rows.append(values)
        finally:
            wb.close()
        return {
            "rows": rows,
            "row_numbers": list(range(row_start, row_end + 1)),
            "col_letters": [get_column_letter(col) for col in range(col_start, col_end + 1)],
            "row_keyword": location.get("row_keyword_address"),
            "col_keyword": location.get("col_keyword_address"),
            "value_address": location.get("value_address"),
            "merged_range": entry["value_meta"].get("merged_range"),
        }
