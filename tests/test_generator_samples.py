"""Sample generator: files exist, open, and carry the intentional messiness."""
from __future__ import annotations

from pathlib import Path

import openpyxl


def test_all_three_sources_exist(input_a: Path) -> None:
    assert (input_a / "prod_harian" / "Prod_14Okt2026.xlsx").is_file()
    assert (input_a / "lap_harian" / "Laporan_Harian_14Okt.xlsx").is_file()
    assert (input_a / "penjualan_harian" / "Penjualan_14Okt.xlsx").is_file()


def test_temp_lock_file_is_present(input_a: Path) -> None:
    assert (input_a / "prod_harian" / "~$Prod_14Okt2026.xlsx").is_file()


def test_workbooks_open_with_expected_sheets(samples_root: Path) -> None:
    prod = openpyxl.load_workbook(samples_root / "input_a" / "prod_harian" / "Prod_14Okt2026.xlsx")
    assert prod.sheetnames == ["Produksi"]
    prod_b = openpyxl.load_workbook(samples_root / "input_b" / "prod_harian" / "Prod_14Okt2026.xlsx")
    assert prod_b.sheetnames == ["Produksi 01-10"]


def test_merged_and_hidden_present_in_prod_sheet(input_a: Path) -> None:
    wb = openpyxl.load_workbook(input_a / "prod_harian" / "Prod_14Okt2026.xlsx")
    ws = wb["Produksi"]
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert "K57:L57" in merged
    assert ws.row_dimensions[50].hidden is True
