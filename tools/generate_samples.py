"""Generate deliberately messy sample Excel files for tests and demos.

One variant = one "pabrik": keyword rows/columns shift, sheet names change, and
numbers differ between variants, mirroring real reports that are not formatted alike.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


@dataclass(frozen=True)
class Layout:
    """How this variant's report is laid out (shifts are in cells)."""

    prod_sheet: str
    lap_sheet: str
    row_shift: int  # data rows in the production sheet
    col_shift: int  # data columns in the production sheet
    pen_shift: int  # data rows in the sales sheet
    prod_value: float
    pen_value: int


#: Variant "a" is the base layout; b/c simulate other factories with shifted data.
LAYOUTS: dict[str, Layout] = {
    "a": Layout("Produksi", "Laporan 01-10", 0, 0, 0, 1523.5, 2_500_000),
    "b": Layout("Produksi 01-10", "Laporan 11-20", 3, 1, 2, 1600.0, 2_750_000),
    "c": Layout("Produksi 11-20", "Laporan 21-31", 5, 2, 4, 1480.25, 2_610_000),
}


def _write_prod(path: Path, lay: Layout) -> None:
    """Production sheet: duplicate keywords, merged cells, formula, hidden row, drift."""
    wb = Workbook()
    ws = wb.active
    ws.title = lay.prod_sheet

    def r(base: int) -> int:
        return base + lay.row_shift

    def c(base: int) -> int:
        return base + lay.col_shift

    ws["A1"] = "LAPORAN PRODUKSI HARIAN"
    # Section banner (merged, does not touch the keyword column I=9).
    ws.cell(row=r(10), column=1, value="SECTION PRODUKSI")
    ws.merge_cells(start_row=r(10), start_column=1, end_row=r(10), end_column=8)
    # Column headers: J=10 Harian, K=11 Akum, L=12 Total (shifted per variant).
    ws.cell(row=r(12), column=c(10), value="Harian")
    ws.cell(row=r(12), column=c(11), value="Akum")
    ws.cell(row=r(12), column=c(12), value="Total")
    # Main row: keyword in I, Harian is a formula that was never calculated.
    ws.cell(row=r(45), column=9, value="Produksi Harian")
    ws.cell(row=r(45), column=c(10), value="=SUM(J13:J44)")
    ws.cell(row=r(45), column=c(11), value=45000.0)
    ws.cell(row=r(45), column=c(12), value=lay.prod_value)
    # Excel error value in a "Total" cell.
    ws.cell(row=r(47), column=9, value="Total Energi")
    ws.cell(row=r(47), column=c(12), value="#REF!")
    # Hidden row.
    ws.cell(row=r(50), column=9, value="Scrap Harian")
    ws.cell(row=r(50), column=c(12), value=12)
    ws.row_dimensions[r(50)].hidden = True
    # Value cell inside a merged range (target L lies in merge K:L, value at K).
    ws.cell(row=r(57), column=9, value="Utilisasi Mesin")
    ws.cell(row=r(57), column=c(11), value=87.2)
    ws.merge_cells(start_row=r(57), start_column=c(11), end_row=r(57), end_column=c(12))
    # Date, text-in-number, and empty values.
    ws.cell(row=r(60), column=9, value="Tanggal Laporan")
    ws.cell(row=r(60), column=c(12), value=datetime(2026, 10, 14))
    ws.cell(row=r(63), column=9, value="Target Harian")
    ws.cell(row=r(63), column=c(12), value="N/A")
    ws.cell(row=r(66), column=9, value="Catatan")
    # Second section with a duplicate of the main keyword.
    ws.cell(row=r(90), column=1, value="SECTION LAINNYA")
    ws.merge_cells(start_row=r(90), start_column=1, end_row=r(90), end_column=8)
    ws.cell(row=r(94), column=9, value="Produksi Harian")
    ws.cell(row=r(94), column=c(12), value=999.0)
    wb.save(path)


def _write_lap(path: Path, lay: Layout) -> None:
    """Daily report sheet: the same keyword in two sections (summary vs production)."""
    wb = Workbook()
    ws = wb.active
    ws.title = lay.lap_sheet
    ws["A1"] = "SUMMARY"
    ws["C5"] = "Total Produksi"
    ws["D5"] = 999
    ws["A20"] = "SECTION PRODUKSI"
    ws["C25"] = "Total Produksi"
    ws["D25"] = lay.prod_value
    wb.save(path)


def _write_pen(path: Path, lay: Layout) -> None:
    """Sales sheet: non-tabular 'label: value' layout, duplicate label, hidden column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Penjualan"

    def r(base: int) -> int:
        return base + lay.pen_shift

    ws.cell(row=r(10), column=1, value="Penjualan Harian")
    ws.cell(row=r(10), column=2, value=lay.pen_value)
    ws.cell(row=r(14), column=1, value="Retur")
    ws.cell(row=r(14), column=2, value=150_000)
    ws.cell(row=r(17), column=1, value="Harga Jual")
    ws.cell(row=r(18), column=1, value=15_000)  # value sits BELOW the label
    ws.cell(row=30, column=2, value="Target")  # value sits LEFT of the label
    ws.cell(row=30, column=1, value=2_400_000)
    ws.cell(row=r(40), column=1, value="Penjualan Harian")  # duplicate label
    ws.cell(row=r(40), column=2, value=1_900_000)
    ws.column_dimensions["B"].hidden = True  # hidden column holding values
    wb.save(path)


def generate(out_root: Path, variant: str) -> Path:
    """Create the sample input tree for one variant; returns the input root.

    Layout: {out_root}/input_{variant}/{prod_harian,lap_harian,penjualan_harian}/...
    """
    lay = LAYOUTS[variant]
    root = out_root / f"input_{variant}"
    prod_dir = root / "prod_harian"
    prod_dir.mkdir(parents=True, exist_ok=True)
    _write_prod(prod_dir / "Prod_14Okt2026.xlsx", lay)
    # Excel lock/temp file that the engine must ignore.
    (prod_dir / "~$Prod_14Okt2026.xlsx").write_bytes(b"temporary lock file")
    lap_dir = root / "lap_harian"
    lap_dir.mkdir(parents=True, exist_ok=True)
    _write_lap(lap_dir / "Laporan_Harian_14Okt.xlsx", lay)
    pen_dir = root / "penjualan_harian"
    pen_dir.mkdir(parents=True, exist_ok=True)
    _write_pen(pen_dir / "Penjualan_14Okt.xlsx", lay)
    return root


def main(argv: list[str] | None = None) -> int:
    """CLI entry: generate one or all sample variants."""
    parser = argparse.ArgumentParser(description="Buat file Excel sampel yang sengaja berantakan.")
    parser.add_argument("--out", default="samples", help="Folder hasil (default: samples)")
    parser.add_argument("--variant", choices=[*LAYOUTS, "all"], default="all")
    args = parser.parse_args(argv)
    out = Path(args.out)
    variants = list(LAYOUTS) if args.variant == "all" else [args.variant]
    for variant in variants:
        print(f"Varian {variant}: {generate(out, variant)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
