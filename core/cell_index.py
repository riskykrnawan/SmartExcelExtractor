"""Workbook loading and the fast per-sheet cell index.

The index maps normalized text -> positions so keyword lookups do not rescan
the sheet for every field. Each sheet is built once per run and cached.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field as dataclass_field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .matcher import fold_text, matches_text, normalize_text
from .models import KeywordConfig, MatchMode

logger = logging.getLogger(__name__)

#: Extensions the engine can read. .xls is intentionally absent (see docs/SPEC.md).
EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
#: Excel lock/temp files must be ignored.
TEMP_PREFIX = "~$"


class ExcelFileError(Exception):
    """A file cannot be opened as an Excel workbook (corrupt, locked, not Excel)."""


def list_excel_files(folder: Path) -> list[Path]:
    """Return Excel files in a folder, newest first; `~$` temp files are skipped.

    Returns an empty list when the folder does not exist or cannot be read.
    """
    try:
        entries = list(folder.iterdir())
    except OSError:
        return []

    def _mtime(path: Path) -> float:
        try:
            return path.stat().st_mtime
        except OSError:
            return 0.0

    files = [
        path
        for path in entries
        if path.is_file()
        and not path.name.startswith(TEMP_PREFIX)
        and path.suffix.lower() in EXCEL_SUFFIXES
    ]
    files.sort(key=lambda path: (-_mtime(path), path.name))
    return files


@dataclass
class ValueRead:
    """Value of a cell with the metadata the audit needs."""

    value: Any
    merged: bool = False
    merged_range: str | None = None
    from_formula: bool = False
    hidden: bool = False


@dataclass
class SheetData:
    """One loaded sheet: values, cell index, merged ranges, hidden rows/columns."""

    title: str
    values: dict[tuple[int, int], Any] = dataclass_field(default_factory=dict)
    texts: dict[tuple[int, int], str] = dataclass_field(default_factory=dict)
    index: dict[str, list[tuple[int, int]]] = dataclass_field(default_factory=dict)
    formulas: dict[tuple[int, int], str] = dataclass_field(default_factory=dict)
    merged_ranges: list[tuple[int, int, int, int]] = dataclass_field(default_factory=list)
    hidden_rows: set[int] = dataclass_field(default_factory=set)
    hidden_cols: set[int] = dataclass_field(default_factory=set)
    max_row: int = 0
    max_col: int = 0

    def find_positions(self, keyword: KeywordConfig) -> list[tuple[int, int]]:
        """All cells matching `keyword`, sorted by (row, col). Empty list when none."""
        if not keyword.text.strip():
            return []
        mode = keyword.match_mode
        if mode is MatchMode.EXACT and not keyword.case_sensitive:
            # Fast path: exact case-insensitive lookup is a single dict hit.
            return list(self.index.get(fold_text(normalize_text(keyword.text)), []))
        positions: list[tuple[int, int]] = []
        if keyword.case_sensitive:
            positions = [pos for pos, text in self.texts.items() if matches_text(text, keyword)]
        else:
            # Index keys are already folded; matches_text folds the keyword itself.
            for key, bucket in self.index.items():
                if matches_text(key, keyword):
                    positions.extend(bucket)
        positions.sort()
        return positions

    def value_at(self, row: int, col: int) -> ValueRead:
        """Read a cell value, resolving merged ranges to their top-left cell.

        Flags `merged` when the target lies anywhere inside a merged range,
        `from_formula` when the cell holds a formula, and `hidden` when the
        original target row/column is hidden.
        """
        hidden = row in self.hidden_rows or col in self.hidden_cols
        target_row, target_col = row, col
        merged = False
        merged_range: str | None = None
        for r1, c1, r2, c2 in self.merged_ranges:
            if r1 <= row <= r2 and c1 <= col <= c2:
                merged = True
                merged_range = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
                target_row, target_col = r1, c1
                break
        return ValueRead(
            value=self.values.get((target_row, target_col)),
            merged=merged,
            merged_range=merged_range,
            from_formula=(target_row, target_col) in self.formulas,
            hidden=hidden,
        )


@dataclass
class LoadedWorkbook:
    """Both views of one file: cached values (data_only) and formulas."""

    path: Path
    values_wb: Any
    formulas_wb: Any

    def sheet_titles(self) -> list[str]:
        """Sheet names in file order."""
        return list(self.values_wb.sheetnames)

    def close(self) -> None:
        """Close both workbook handles."""
        self.values_wb.close()
        self.formulas_wb.close()


def load_workbook_pair(path: Path) -> LoadedWorkbook:
    """Open a file twice: values (data_only=True) and formulas (data_only=False).

    The second pass is required for `from_formula` and FORMULA_NOT_CALCULATED.
    Raises ExcelFileError with a user-friendly message on any failure.
    """
    try:
        values_wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any failure means the file is unusable
        logger.exception("Cannot open %s", path)
        raise ExcelFileError(
            f"File '{path.name}' tidak bisa dibuka (rusak, terkunci, atau bukan file Excel)."
        ) from exc
    try:
        formulas_wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:  # noqa: BLE001
        values_wb.close()
        logger.exception("Cannot open formulas view of %s", path)
        raise ExcelFileError(
            f"File '{path.name}' tidak bisa dibuka (rusak, terkunci, atau bukan file Excel)."
        ) from exc
    return LoadedWorkbook(path=path, values_wb=values_wb, formulas_wb=formulas_wb)


def build_sheet_data(ws_values: Any, ws_formulas: Any, title: str) -> SheetData:
    """Build the cell index and metadata for one sheet (called once per run)."""
    data = SheetData(title=title)
    for row in ws_values.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            pos = (cell.row, cell.column)
            data.values[pos] = cell.value
            text = normalize_text(cell.value)
            if text:
                data.texts[pos] = text
                data.index.setdefault(fold_text(text), []).append(pos)
    for bucket in data.index.values():
        bucket.sort()
    for rng in ws_values.merged_cells.ranges:
        data.merged_ranges.append((rng.min_row, rng.min_col, rng.max_row, rng.max_col))
    for row_idx, dim in ws_values.row_dimensions.items():
        if dim.hidden:
            data.hidden_rows.add(row_idx)
    for letter, dim in ws_values.column_dimensions.items():
        if dim.hidden:
            try:
                data.hidden_cols.add(column_index_from_string(str(letter)))
            except ValueError:
                logger.warning("Unexpected column dimension key %r in sheet %r", letter, title)
    for row in ws_formulas.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                data.formulas[(cell.row, cell.column)] = value
    data.max_row = ws_values.max_row or 0
    data.max_col = ws_values.max_column or 0
    return data
